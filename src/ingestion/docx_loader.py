"""Loader implementation for DOCX Word documents using python-docx."""

from pathlib import Path
from typing import Union
import docx
from src.domain.documents import Document, DocumentMetadata, DocumentType, ParsedPage
from src.utils.file_utils import compute_file_hash
from src.utils.text_utils import normalize_whitespace


class DocxLoader:
    """Extracts text, paragraphs, and metadata from DOCX files."""

    def load(self, file_path: Union[str, Path]) -> Document:
        path = Path(file_path).resolve()
        if not path.is_file():
            raise FileNotFoundError(f"DOCX file does not exist: {path}")

        file_hash = compute_file_hash(path)
        file_size = path.stat().st_size
        doc_id = f"doc_{file_hash[:16]}"

        doc = docx.Document(str(path))

        paragraphs_text: list[str] = []
        for p in doc.paragraphs:
            cleaned_line = normalize_whitespace(p.text)
            if cleaned_line:
                paragraphs_text.append(cleaned_line)

        # Also extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_text = [normalize_whitespace(cell.text) for cell in row.cells if cell.text]
                if row_text:
                    paragraphs_text.append(" | ".join(row_text))

        concatenated_text = "\n\n".join(paragraphs_text)

        # In Word documents, pages are flow-based; we simulate logical pages per paragraphs or treat as page 1
        parsed_pages = [
            ParsedPage(
                page_number=1,
                text=concatenated_text,
                metadata={"source_page": 1, "doc_id": doc_id},
            )
        ]

        title = None
        if doc.core_properties and doc.core_properties.title:
            title = doc.core_properties.title.strip()
        if not title:
            title = path.stem.replace("_", " ").title()

        metadata = DocumentMetadata(
            source_path=str(path),
            file_name=path.name,
            file_type=DocumentType.DOCX,
            file_size_bytes=file_size,
            file_hash=file_hash,
            total_pages=1,
            title=title,
        )

        return Document(
            doc_id=doc_id,
            content=concatenated_text,
            pages=parsed_pages,
            metadata=metadata,
        )